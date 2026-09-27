"""Designer-facing Excel model. Legacy evidence remains in the audit viewer."""

from __future__ import annotations

import json
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
BOOK = ROOT / "data/unified-quest/UnifiedQuest.xlsx"
EXPORT_DIR = ROOT / "outputs/unified-quest/export"

SCHEMA = {
    "EventType": [
        ("event_key", "事件编号"), ("name", "事件名称"), ("category", "玩法分类"),
        ("trigger_point", "何时触发"), ("counter_mode", "如何计数"), ("unit", "计数单位"),
        ("scope", "归属范围"), ("description", "策划说明"),
        ("steam_scope", "Steam范围"), ("scope_reason", "范围依据"),
        ("publish_state", "定义状态"),
    ],
    "EventParam": [
        ("event_key", "所属事件"), ("param_key", "参数编号"), ("name", "参数名称"),
        ("data_type", "值类型"), ("param_kind", "参数用途"),
        ("allowed_values", "可配置范围"), ("comparison", "筛选方式"),
        ("example", "填写示例"), ("description", "参数说明"),
    ],
    "QuestType": [
        ("type_key", "任务类型编号"), ("name", "任务类型名称"),
        ("structure", "默认规则"), ("description", "类型说明"),
        ("publish_state", "发布状态"),
    ],
    "Quest": [
        ("quest_id", "任务编号"), ("name", "任务名称"), ("quest_type", "任务分类"),
        ("legacy_group", "旧列表分组"),
        ("scope", "归属范围"), ("repeat_cycle", "重复周期"), ("entry", "打开入口"),
        ("min_level", "开放等级"), ("pre_quest_id", "前置任务"),
        ("reward_ref", "奖励引用"), ("publish_state", "发布状态"),
    ],
    "QuestCondition": [
        ("condition_id", "条件编号"), ("quest_id", "所属任务"), ("phase", "判定阶段"),
        ("event_key", "事件类型"), ("filter_param", "筛选参数"),
        ("filter_operator", "筛选比较"), ("filter_value", "筛选值"),
        ("operator", "目标比较"), ("target_value", "目标数量"),
        ("logic_group", "条件组合"), ("count_start", "计数起点"), ("scope", "归属范围"),
    ],
    "QuestNode": [
        ("node_id", "节点编号"), ("quest_id", "所属任务"), ("node_type", "节点类型"),
        ("reference", "条件或对话组ID"), ("next_node", "下一节点"), ("entry", "进入方式"),
        ("dialogue_phase", "对话状态"),
    ],
    "DialogueLine": [
        ("sequence_id", "对话组ID"), ("line_order", "行序"), ("speaker", "说话人"),
        ("position", "立绘位置"), ("text", "对白内容"), ("scale", "缩放"),
        ("speed", "播放速度"), ("delay", "延时"), ("show_skip", "显示跳过"),
    ],
    "LegacyMap": [
        ("source", "旧来源"), ("legacy_id", "旧编号"), ("record_kind", "记录类型"),
        ("new_id", "统一编号"), ("migration_state", "核对状态"), ("note", "证据与说明"),
    ],
}

CONFIRMED_EVENTS = [
    {
        "event_key": "EMQCT_2", "name": "神将招募完成", "category": "招募",
        "trigger_point": "单抽或十连结算并获得神将后触发",
        "counter_mode": "累加实际招募次数；单抽加 1，十连加 10", "unit": "次", "scope": "角色",
        "description": "三种招募类型共用；次数由服务端实际结算结果决定。", "publish_state": "草稿",
    },
    {
        "event_key": "EMQCT_4", "name": "角色登录完成", "category": "登录",
        "trigger_point": "选择角色并完成服务端登录流程后触发",
        "counter_mode": "每次成功进入角色加 1", "unit": "次", "scope": "角色",
        "description": "每日登录任务按当日周期统计，重复登录不会把目标 1 重复完成。", "publish_state": "草稿",
    },
    {
        "event_key": "EMQCT_7", "name": "装备强化完成", "category": "装备",
        "trigger_point": "强化消耗与属性更新完成后触发；批量强化按实际完成数量触发",
        "counter_mode": "累加实际强化操作数", "unit": "次", "scope": "角色",
        "description": "单件强化加 1；批量强化使用服务端 qhCnt 计数。", "publish_state": "草稿",
    },
    {
        "event_key": "EMQCT_8", "name": "装备精炼完成", "category": "装备",
        "trigger_point": "精炼消耗与属性更新完成后触发",
        "counter_mode": "每次成功精炼操作加 1", "unit": "次", "scope": "角色",
        "description": "原任务计成功操作次数，不按提升的精炼等级计数。", "publish_state": "草稿",
    },
    {
        "event_key": "EMQCT_11", "name": "法宝强化完成", "category": "法宝",
        "trigger_point": "法宝强化消耗、等级与属性更新完成后触发",
        "counter_mode": "每次成功强化操作加 1", "unit": "次", "scope": "角色",
        "description": "原任务计成功操作次数，不按一次提升的等级计数。", "publish_state": "草稿",
    },
    {
        "event_key": "EMQCT_12", "name": "封神列传战斗胜利", "category": "封神列传",
        "trigger_point": "列传战斗结果为胜利并推进关卡后触发",
        "counter_mode": "每场胜利加 1", "unit": "场", "scope": "角色",
        "description": "失败不计数；同一胜利只发布一次。", "publish_state": "草稿",
    },
    {
        "event_key": "EMQCT_14", "name": "法宝搜索完成", "category": "法宝",
        "trigger_point": "搜索结算、扣除搜索次数并发放产出后；手动与自动搜索均触发",
        "counter_mode": "累加实际搜索次数", "unit": "次", "scope": "角色",
        "description": "按实际执行次数增加进度；打开法宝界面或领取任务均不计数。",
        "publish_state": "草稿",
    },
    {
        "event_key": "EMQCT_34", "name": "法宝合成完成", "category": "法宝",
        "trigger_point": "消耗材料并生成法宝后；手动与自动合成均触发",
        "counter_mode": "累加实际合成数量", "unit": "件", "scope": "角色",
        "description": "产物品质来自 fabao.json；条件可按最低品质筛选，奖励差异不拆事件。",
        "publish_state": "草稿",
    },
    {
        "event_key": "EMQCT_39", "name": "当日活跃度更新", "category": "每日任务",
        "trigger_point": "服务端增加当日活跃度并写入角色状态后触发",
        "counter_mode": "采用更新后的当日活跃度当前值，不累加事件次数", "unit": "点", "scope": "角色",
        "description": "任务条件比较当日活跃度与目标值；跨日归零由原每日状态管理。", "publish_state": "草稿",
    },
]

CONFIRMED_PARAMS = [
    {"event_key": "EMQCT_2", "param_key": "count", "name": "实际招募次数",
     "data_type": "整数", "param_kind": "计数增量", "allowed_values": "1 或 10",
     "comparison": "累加", "example": 10, "description": "单抽为 1，十连为 10。"},
    {"event_key": "EMQCT_4", "param_key": "count", "name": "本次登录",
     "data_type": "整数", "param_kind": "计数增量", "allowed_values": "1",
     "comparison": "累加", "example": 1, "description": "角色登录流程完成时加 1。"},
    {"event_key": "EMQCT_7", "param_key": "count", "name": "实际强化操作数",
     "data_type": "整数", "param_kind": "计数增量", "allowed_values": "大于 0",
     "comparison": "累加", "example": 1, "description": "单件为 1；批量取服务端实际计数。"},
    {"event_key": "EMQCT_8", "param_key": "count", "name": "本次精炼操作",
     "data_type": "整数", "param_kind": "计数增量", "allowed_values": "1",
     "comparison": "累加", "example": 1, "description": "每次成功精炼操作加 1。"},
    {"event_key": "EMQCT_11", "param_key": "count", "name": "本次强化操作",
     "data_type": "整数", "param_kind": "计数增量", "allowed_values": "1",
     "comparison": "累加", "example": 1, "description": "每次成功法宝强化操作加 1。"},
    {"event_key": "EMQCT_12", "param_key": "count", "name": "胜利场数",
     "data_type": "整数", "param_kind": "计数增量", "allowed_values": "1",
     "comparison": "累加", "example": 1, "description": "列传战斗胜利后加 1。"},
    {
        "event_key": "EMQCT_14", "param_key": "count", "name": "实际搜索次数",
        "data_type": "整数", "param_kind": "计数增量", "allowed_values": "大于 0",
        "comparison": "累加", "example": 1,
        "description": "手动或自动搜索时，本次实际消耗并完成的搜索次数。",
    },
    {
        "event_key": "EMQCT_34", "param_key": "count", "name": "实际合成数量",
        "data_type": "整数", "param_kind": "计数增量", "allowed_values": "大于 0",
        "comparison": "累加", "example": 1,
        "description": "本次成功生成的法宝件数；自动合成按实际件数。",
    },
    {
        "event_key": "EMQCT_34", "param_key": "quality", "name": "产物品质",
        "data_type": "整数", "param_kind": "筛选参数", "allowed_values": "fabao.json 的 quality：3、4、5、6、7",
        "comparison": "大于等于", "example": 3,
        "description": "只统计产物品质达到配置下限的合成数量。",
    },
    {"event_key": "EMQCT_39", "param_key": "current_value", "name": "当日活跃度当前值",
     "data_type": "整数", "param_kind": "状态值", "allowed_values": "大于等于 0",
     "comparison": "取当前值", "example": 50, "description": "服务端写入后的当日累计活跃度。"},
]

LEGACY_EVENT_KEYS = {key: key for key in (
    "EMQCT_2", "EMQCT_4", "EMQCT_7", "EMQCT_8", "EMQCT_11",
    "EMQCT_12", "EMQCT_14", "EMQCT_34", "EMQCT_39",
)}

QUEST_TYPES = [
    {"type_key": "main", "name": "主线", "structure": "一次性推进；按前置任务解锁",
     "description": "主线任务可包含对白、事件目标和状态检查节点。", "publish_state": "草稿"},
    {"type_key": "side", "name": "支线", "structure": "一次性推进；满足开放条件后接取",
     "description": "支线任务可包含对白和事件目标；当前旧支线数据仅作迁移核对。", "publish_state": "草稿"},
    {"type_key": "daily", "name": "日常", "structure": "每日重置任务进度",
     "description": "活跃奖励与玩法阶梯是任务包或展示组，不另建任务类型。", "publish_state": "草稿"},
    {"type_key": "cumulative", "name": "累计", "structure": "跨日累计；可串接阶段目标",
     "description": "目标进度不随每日刷新清零；法宝搜索累计目标属于此类。", "publish_state": "草稿"},
]

DIALOGUE_PHASES = (
    ("before_accept", "未接取前对话"),
    ("accepted_incomplete", "已接未完成对话"),
    ("completed_unclaimed", "已完成未领取时对话"),
)
DIALOGUE_PHASE_KEYS = {key for key, _ in DIALOGUE_PHASES}


def add_fixed_dialogue_slots(records, quest_id):
    """Each task has three state-bound dialogue slots, even if their group IDs are empty."""
    existing = {row.get("dialogue_phase") for row in records["QuestNode"]
                if str(row.get("quest_id")) == str(quest_id) and row.get("node_type") == "Dialogue"}
    for phase, _ in DIALOGUE_PHASES:
        if phase in existing:
            continue
        records["QuestNode"].append({"node_id": f"{quest_id}.dialogue.{phase}",
            "quest_id": quest_id, "node_type": "Dialogue", "reference": "",
            "next_node": "", "entry": "", "dialogue_phase": phase})


def source_dialogue_lines(source_rows):
    """Import the current chapter dialogue groups without guessing task associations."""
    result, seen = [], {}
    for source in source_rows:
        group = str(int(source["dialogid"]))
        order = int(source["order"])
        key = (group, order)
        values = (str(source.get("comment") or source.get("npcid") or ""),
                  int(source.get("position") or 0), str(source.get("dialog") or ""),
                  int(source.get("scale") or 1), int(source.get("speed") or 0),
                  int(source.get("delay") or 0), int(source.get("showskip") or 0))
        if key in seen:
            if seen[key] != values:
                raise ValueError(f"对话组 {group} 第 {order} 句有冲突的源记录")
            continue
        seen[key] = values
        result.append(dict(sequence_id=group, line_order=order, speaker=values[0],
            position=values[1], text=values[2], scale=values[3], speed=values[4],
            delay=values[5], show_skip=values[6]))
    return result

# Steam module scope is separate from whether an event's parameter contract is complete.
# The legacy EventType row remains available for migration audit even when hidden here.
STEAM_EXCLUDED_EVENTS = {
    3: "血战到底", 5: "好友赠送", 6: "帮派", 9: "竞技场",
    17: "支线副本入口", 33: "血战到底", 37: "七日目标",
    38: "充值活动", 43: "竞技场", 44: "血战到底",
    54: "七日目标", 55: "封神试炼入口", 56: "决战昆仑",
    59: "帮派", 60: "帮派副本", 61: "基金",
}
STEAM_UNVERIFIED_EVENTS = {
    1: "商城购买入口未核对", 18: "商店刷新入口未核对",
    31: "VIP入口未核对", 36: "旧寻宝与法宝搜索的关系未核对",
}


def steam_scope(event_key):
    number = int(str(event_key).split("_")[-1])
    if number in STEAM_EXCLUDED_EVENTS:
        return "屏蔽", STEAM_EXCLUDED_EVENTS[number]
    if number in STEAM_UNVERIFIED_EVENTS:
        return "待核对", STEAM_UNVERIFIED_EVENTS[number]
    return "纳入", "当前保留玩法或通用角色状态"


def clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def read_book(path: Path = BOOK):
    book = load_workbook(path, read_only=True, data_only=True)
    result = {}
    for table, columns in SCHEMA.items():
        page = book[table]
        expected = [key for key, _ in columns]
        actual = [cell.value for cell in page[2]][:len(expected)]
        if actual != expected:
            raise ValueError(f"{table} 字段定义与策划模型不符")
        result[table] = [dict(zip(expected, map(clean, values))) for values in
                         page.iter_rows(min_row=3, max_col=len(expected), values_only=True)
                         if any(value is not None and str(value).strip() for value in values)]
    book.close()
    return result


def seed(snapshot, previous=None):
    """Deduplicate EMQCT definitions and daily condition types into one event catalog."""
    previous = previous or {}
    old_quests = {str(row["quest_id"]): row for row in previous.get("Quest", [])}
    old_conditions = {str(row["condition_id"]): row for row in previous.get("QuestCondition", [])}
    old_nodes = defaultdict(list)
    for row in previous.get("QuestNode", []):
        old_nodes[str(row["quest_id"])].append(row)
    old_events = {str(row["event_key"]): row for row in previous.get("EventType", [])}
    records = {table: [] for table in SCHEMA}
    confirmed = {row["event_key"]: row for row in CONFIRMED_EVENTS}
    sources = {source["name"]: source["records"] for source in snapshot["sources"]}
    daily_usage = defaultdict(list)
    for source in sources["每日任务源表"]:
        parts = [part.strip() for part in str(source.get("condition", "")).split(",")]
        if parts and parts[0].isdigit():
            daily_usage[f"EMQCT_{int(parts[0])}"].append(source)
    emqct = {event["legacy_key"]: event for event in snapshot["event_candidates"]
             if event["namespace"] == "EMQCT"}
    keys = set(emqct) | set(daily_usage)
    for key in sorted(keys, key=lambda value: int(value.split("_")[-1])):
        if key in confirmed:
            row = dict(confirmed[key])
            if old_events.get(key, {}).get("publish_state") == "已验收":
                row["publish_state"] = "已验收"
        else:
            event = emqct.get(key, {})
            number = int(key.split("_")[-1])
            uses = daily_usage.get(key, [])
            row = {"event_key": key, "name": event.get("description") or str(uses[0].get("des", key) if uses else key),
                   "category": "", "trigger_point": "", "counter_mode": "", "unit": "", "scope": "",
                   "description": f"现行 daily 引用 {len(uses)} 条；定义来源为 EMQCT。" if uses else "定义来源为 EMQCT；现行 daily 未引用。",
                   "publish_state": "待定义"}
        row["steam_scope"], row["scope_reason"] = steam_scope(key)
        if row["publish_state"] == "Steam排除":
            row["publish_state"] = "待定义"
        records["EventType"].append(row)
    records["EventParam"] = [dict(row) for row in CONFIRMED_PARAMS]
    records["QuestType"] = [dict(row) for row in QUEST_TYPES]
    records["DialogueLine"] = [dict(row) for row in previous.get("DialogueLine", [])] or source_dialogue_lines(
        sources["章节剧情源表"])
    active_quests = set()
    for source in sources["每日任务源表"]:
        parts = [int(part.strip()) for part in str(source.get("condition", "")).split(",") if part.strip()]
        legacy_key = f"EMQCT_{parts[0]}" if parts else ""
        if legacy_key not in LEGACY_EVENT_KEYS:
            continue
        quest_id = f"daily.{int(source['id'])}"
        active_quests.add(quest_id)
        prior = old_quests.get(quest_id, {})
        records["Quest"].append({
            "quest_id": quest_id, "name": prior.get("name") or str(source.get("des", "")),
            "quest_type": "daily" if source.get("daily") else "cumulative",
            "legacy_group": int(source.get("type", 0)),
            "scope": "角色", "repeat_cycle": "每日" if source.get("daily") else "累计",
            "entry": "法宝搜索任务页" if int(source.get("type", 0)) == 3 else "每日任务页",
            "min_level": int(source.get("level") or 0),
            "pre_quest_id": f"daily.{int(source['show'])}" if source.get("show") else "",
            "reward_ref": prior.get("reward_ref") or str(source.get("reward", "")),
            "publish_state": "已验收" if prior.get("migration_state") == "已验收" else "草稿",
        })
        old_condition = old_conditions.get(f"{quest_id}.complete", {})
        records["QuestCondition"].append({
            "condition_id": f"{quest_id}.complete", "quest_id": quest_id, "phase": "完成",
            "event_key": LEGACY_EVENT_KEYS[legacy_key], "filter_param": "",
            "filter_operator": "", "filter_value": "", "operator": ">=",
            "target_value": old_condition.get("target_value", parts[1] if len(parts) > 1 else 1),
            "logic_group": "AND", "count_start": "周期开始" if source.get("daily") else "角色创建",
            "scope": "角色",
        })
        if old_nodes.get(quest_id):
            records["QuestNode"].extend(dict(node, dialogue_phase=node.get("dialogue_phase", ""))
                                        for node in old_nodes[quest_id])
        else:
            for kind, suffix, next_suffix, reference in (("Start", "start", "wait", ""),
                ("WaitEvent", "wait", "complete", f"{quest_id}.complete"),
                ("Complete", "complete", "", "")):
                records["QuestNode"].append({
                    "node_id": f"{quest_id}.{suffix}", "quest_id": quest_id, "node_type": kind,
                    "reference": reference, "next_node": f"{quest_id}.{next_suffix}" if next_suffix else "",
                    "entry": "任务列表" if kind == "Start" else "", "dialogue_phase": "",
                })
        add_fixed_dialogue_slots(records, quest_id)
    for event in snapshot["event_candidates"]:
        old_key = event["legacy_key"]
        target = old_key if event["namespace"] == "EMQCT" else ""
        if target in confirmed:
            note = f"{event['description']}；已核对事件参数与产生点。"
        elif target:
            note = f"{event['description']}；事件类型已列入汇总，参数与产生点待核对。"
        else:
            note = f"{event['description'] or old_key}；旧条件/事件候选，语义未归并。"
        records["LegacyMap"].append({"source": event["namespace"], "legacy_id": event["id"],
            "record_kind": "事件", "new_id": target,
            "migration_state": "已定义待验收" if target in confirmed else "待核对", "note": note})
    for source in sources["每日任务源表"]:
        quest_id = f"daily.{int(source['id'])}"
        target = quest_id if quest_id in active_quests else ""
        records["LegacyMap"].append({"source": "daily.xlsx", "legacy_id": int(source["id"]),
            "record_kind": "任务", "new_id": target,
            "migration_state": "已导入草稿" if target else "待核对",
            "note": "已导入策划任务表" if target else "旧任务配置可在审计页查看，事件合同未确认"})
    story_groups = defaultdict(list)
    for row in sources["章节剧情源表"]:
        story_groups[int(row["dialogid"])].append(row)
    for number, lines in sorted(story_groups.items()):
        note = "剧情源表已解析，Steam 播放入口待核对；原文在审计页。"
        if number == 10042:
            note += "第 37、39 行完全重复。"
        records["LegacyMap"].append({"source": "mission_dialog.xlsx", "legacy_id": number,
            "record_kind": "剧情", "new_id": str(number), "migration_state": "已导入对话组", "note": note})
    for row in sources["旧任务源表"]:
        records["LegacyMap"].append({"source": "mission_config.xlsx", "legacy_id": clean(row["id"]),
            "record_kind": "任务", "new_id": "", "migration_state": "历史休眠",
            "note": "旧任务加载入口已注释，原配置在审计页。"})
    return records


def validate(records):
    errors = []
    for table, columns in SCHEMA.items():
        if table not in records or not isinstance(records[table], list):
            errors.append(f"缺少表 {table}")
            continue
        allowed = {key for key, _ in columns}
        for index, row in enumerate(records[table], 3):
            if not isinstance(row, dict) or set(row) - allowed:
                errors.append(f"{table} 第 {index} 行字段不合法")
    if errors:
        return errors
    def unique(table, key):
        seen = set()
        for index, row in enumerate(records[table], 3):
            value = str(row.get(key, ""))
            if not value:
                errors.append(f"{table} 第 {index} 行缺少 {key}")
            elif value in seen:
                errors.append(f"{table} 编号重复：{value}")
            seen.add(value)
        return seen
    events = unique("EventType", "event_key")
    scoped_events = {row["event_key"] for row in records["EventType"]
                     if row.get("steam_scope") == "纳入"}
    selectable = {row["event_key"] for row in records["EventType"]
                  if row.get("steam_scope") == "纳入" and row.get("publish_state") in ("草稿", "已验收")}
    quests = unique("Quest", "quest_id")
    conditions = unique("QuestCondition", "condition_id")
    nodes = unique("QuestNode", "node_id")
    params = {}
    for row in records["EventType"]:
        key = row["event_key"]
        if not row.get("name"):
            errors.append(f"事件 {key} 缺少名称")
        if row.get("publish_state") not in ("待定义", "草稿", "已验收"):
            errors.append(f"事件 {key} 发布状态无效")
        if row.get("steam_scope") not in ("纳入", "屏蔽", "待核对") or not row.get("scope_reason"):
            errors.append(f"事件 {key} 缺少 Steam 范围依据")
        if key in selectable and any(not str(row.get(field, "")).strip() for field in
               ("category", "trigger_point", "counter_mode", "unit", "scope", "description")):
            errors.append(f"事件 {key} 缺少可配置的业务定义")
    for row in records["EventParam"]:
        key = (str(row.get("event_key", "")), str(row.get("param_key", "")))
        if key in params:
            errors.append(f"事件参数重复：{key[0]} / {key[1]}")
        params[key] = row
        if key[0] not in events or any(not str(row.get(field, "")).strip() for field in
                ("param_key", "name", "data_type", "param_kind", "allowed_values", "comparison", "example", "description")):
            errors.append(f"事件参数 {key[0]} / {key[1]} 定义不完整")
        if row.get("param_kind") not in ("计数增量", "筛选参数", "状态值"):
            errors.append(f"事件参数 {key[0]} / {key[1]} 用途无效")
    for event_key in selectable:
        if not any(key[0] == event_key and param.get("param_kind") in ("计数增量", "状态值")
                   for key, param in params.items()):
            errors.append(f"事件 {event_key} 缺少计数或状态参数")
    quest_types = unique("QuestType", "type_key")
    if set(quest_types) != {row["type_key"] for row in QUEST_TYPES}:
        errors.append("任务类型必须且只能是主线、支线、日常、累计")
    for row in records["QuestType"]:
        if any(not str(row.get(field, "")).strip() for field in ("name", "structure", "description")):
            errors.append(f"任务类型 {row['type_key']} 定义不完整")
        if row.get("publish_state") not in ("草稿", "已验收"):
            errors.append(f"任务类型 {row['type_key']} 发布状态无效")
    quest_rows = {str(row.get("quest_id", "")): row for row in records["Quest"]}
    for row in records["Quest"]:
        if row.get("quest_type") not in quest_types:
            errors.append(f"任务 {row['quest_id']} 引用了未定义的任务类型")
        expected_cycle = {"main": "一次", "side": "一次", "daily": "每日", "cumulative": "累计"}.get(row.get("quest_type"))
        if expected_cycle and row.get("repeat_cycle") != expected_cycle:
            errors.append(f"任务 {row['quest_id']} 的周期与任务类型不一致")
        if row.get("pre_quest_id") and str(row["pre_quest_id"]) not in quests:
            errors.append(f"任务 {row['quest_id']} 前置任务不存在")
        if row.get("publish_state") not in ("草稿", "已验收"):
            errors.append(f"任务 {row['quest_id']} 发布状态无效")
    for row in records["QuestCondition"]:
        cid = row.get("condition_id")
        if str(row.get("quest_id", "")) not in quests:
            errors.append(f"条件 {cid} 引用不存在的任务")
        event_key = str(row.get("event_key", ""))
        if event_key not in scoped_events:
            errors.append(f"条件 {cid} 引用的事件不在当前 Steam 范围")
        elif quest_rows.get(str(row.get("quest_id", "")), {}).get("publish_state") == "已验收":
            if not any(event["event_key"] == event_key and event["publish_state"] == "已验收"
                       for event in records["EventType"]):
                errors.append(f"条件 {cid} 引用的事件尚未验收，任务不能标为已验收")
        if row.get("operator") not in (">=", "=", "<=", ">", "<"):
            errors.append(f"条件 {cid} 目标比较方式无效")
        if not isinstance(row.get("target_value"), (int, float)) or row["target_value"] <= 0:
            errors.append(f"条件 {cid} 目标数量必须大于 0")
        filter_param = str(row.get("filter_param", ""))
        if filter_param:
            parameter = params.get((str(row.get("event_key")), filter_param))
            if parameter is None or parameter.get("param_kind") != "筛选参数":
                errors.append(f"条件 {cid} 筛选参数没有定义")
            if row.get("filter_operator") not in (">=", "=", "<=", ">", "<"):
                errors.append(f"条件 {cid} 筛选比较方式无效")
            if row.get("filter_value", "") == "":
                errors.append(f"条件 {cid} 缺少筛选值")
            elif parameter and parameter.get("data_type") == "整数":
                value = row["filter_value"]
                if isinstance(value, bool) or not isinstance(value, (int, float)) or int(value) != value:
                    errors.append(f"条件 {cid} 筛选值必须是整数")
                elif row.get("event_key") == "EMQCT_34" and filter_param == "quality" and value not in (3, 4, 5, 6, 7):
                    errors.append(f"条件 {cid} 产物品质只能填写 3、4、5、6、7")
        elif row.get("filter_operator") or row.get("filter_value"):
            errors.append(f"条件 {cid} 筛选值未指定参数")
    dialogue_keys, dialogue = set(), set()
    for row in records["DialogueLine"]:
        key = (str(row.get("sequence_id", "")), str(row.get("line_order", "")))
        if not all(key) or key in dialogue_keys:
            errors.append(f"对白行序缺失或重复：{key[0]} / {key[1]}")
        dialogue_keys.add(key)
        dialogue.add(key[0])
        if not isinstance(row.get("line_order"), int) or row["line_order"] < 1:
            errors.append(f"对话组 {key[0]} 的行序必须是正整数")
        if not str(row.get("speaker", "")).strip() or not str(row.get("text", "")).strip():
            errors.append(f"对白 {key[0]} / {key[1]} 缺少说话人或正文")
    by_node = {str(row["node_id"]): row for row in records["QuestNode"]}
    starts = defaultdict(list)
    dialogue_slots = defaultdict(set)
    for row in records["QuestNode"]:
        qid, nid = str(row.get("quest_id", "")), str(row.get("node_id", ""))
        if qid not in quests:
            errors.append(f"节点 {nid} 引用不存在的任务")
        if row.get("node_type") == "Start":
            starts[qid].append(nid)
        if row.get("node_type") not in ("Start", "Dialogue", "WaitEvent", "CheckState", "Complete"):
            errors.append(f"节点 {nid} 类型无效")
        phase = str(row.get("dialogue_phase") or "")
        if row.get("node_type") == "Dialogue":
            if phase not in DIALOGUE_PHASE_KEYS or phase in dialogue_slots[qid]:
                errors.append(f"任务 {qid} 对话状态缺失、重复或无效：{phase}")
            dialogue_slots[qid].add(phase)
            if nid != f"{qid}.dialogue.{phase}" or row.get("next_node") or row.get("entry"):
                errors.append(f"任务 {qid} 的固定对话节点 {nid} 结构不正确")
            if row.get("reference") and str(row["reference"]) not in dialogue:
                errors.append(f"节点 {nid} 引用不存在的对话组")
        elif phase:
            errors.append(f"节点 {nid} 不应填写对话状态")
        next_id = str(row.get("next_node") or "")
        if next_id and next_id not in nodes:
            errors.append(f"节点 {nid} 下一节点不存在")
        if next_id in by_node and str(by_node[next_id].get("quest_id")) != qid:
            errors.append(f"节点 {nid} 跨任务连接")
        if row.get("node_type") == "WaitEvent" and str(row.get("reference", "")) not in conditions:
            errors.append(f"节点 {nid} 条件引用不存在")
    for qid in quests:
        if dialogue_slots[qid] != DIALOGUE_PHASE_KEYS:
            errors.append(f"任务 {qid} 必须配置三个固定对话状态节点")
        if len(starts[qid]) != 1:
            errors.append(f"任务 {qid} 应有一个开始节点")
            continue
        visited, current = set(), starts[qid][0]
        while current:
            if current in visited:
                errors.append(f"任务 {qid} 流程存在循环")
                break
            visited.add(current)
            current = str(by_node[current].get("next_node") or "")
        if not any(by_node[nid].get("node_type") == "Complete" for nid in visited):
            errors.append(f"任务 {qid} 无可达完成节点")
        for row in records["QuestNode"]:
            if row.get("quest_id") == qid and row.get("node_type") != "Dialogue" and str(row.get("node_id")) not in visited:
                errors.append(f"任务 {qid} 存在不可达节点 {row.get('node_id')}")
    legacy_keys = set()
    for row in records["LegacyMap"]:
        key = (str(row.get("source", "")), str(row.get("legacy_id", "")), str(row.get("record_kind", "")))
        if key in legacy_keys:
            errors.append(f"旧编号映射重复：{' / '.join(key)}")
        legacy_keys.add(key)
        target = str(row.get("new_id", ""))
        target_pool = events if row.get("record_kind") == "事件" else dialogue if row.get("record_kind") == "剧情" else quests
        if target and target not in target_pool:
            errors.append(f"旧编号映射 {key[0]} / {key[1]} 的统一编号不存在")
    return errors


def export(records, directory: Path = EXPORT_DIR, preview: bool = False):
    errors = validate(records)
    if errors:
        raise ValueError("；".join(errors[:12]))
    directory.mkdir(parents=True, exist_ok=True)
    allowed = {"草稿", "已验收"} if preview else {"已验收"}
    events = [row for row in records["EventType"]
              if row["steam_scope"] == "纳入" and (preview or row["publish_state"] in allowed)]
    event_keys = {row["event_key"] for row in events}
    quests = [row for row in records["Quest"] if row["publish_state"] in allowed]
    quest_ids = {row["quest_id"] for row in quests}
    quest_type_keys = {row["quest_type"] for row in quests}
    quest_types = [row for row in records["QuestType"] if row["type_key"] in quest_type_keys]
    conditions = [row for row in records["QuestCondition"] if row["quest_id"] in quest_ids]
    for row in conditions:
        if row["event_key"] not in event_keys:
            raise ValueError(f"任务 {row['quest_id']} 的事件尚未允许导出")
    nodes = [row for row in records["QuestNode"] if row["quest_id"] in quest_ids]
    params = [row for row in records["EventParam"] if row["event_key"] in event_keys]
    dialogue_slots = [dict(quest_id=row["quest_id"], phase=row["dialogue_phase"],
                           sequence_id=str(row["reference"] or "")) for row in nodes
                      if row["node_type"] == "Dialogue"]
    dialogue_ids = {row["sequence_id"] for row in dialogue_slots if row["sequence_id"]}
    dialogue_lines = [row for row in records["DialogueLine"] if str(row["sequence_id"]) in dialogue_ids]
    server = {"schema_version": 4, "preview_only": preview, "events": events, "event_params": params,
              "quest_types": quest_types,
               "quests": quests, "conditions": conditions, "nodes": nodes,
               "dialogue_slots": dialogue_slots, "dialogue": dialogue_lines}
    client = {"schema_version": 4, "preview_only": preview,
        "quest_types": [dict(type_key=row["type_key"], name=row["name"], structure=row["structure"])
                        for row in quest_types],
        "quests": [dict(quest_id=row["quest_id"], name=row["name"], quest_type=row["quest_type"],
                        entry=row["entry"], repeat_cycle=row["repeat_cycle"]) for row in quests],
        "events": [dict(event_key=row["event_key"], name=row["name"], unit=row["unit"])
                   for row in events],
        "conditions": [dict(quest_id=row["quest_id"], event_key=row["event_key"],
                            filter_param=row["filter_param"], filter_operator=row["filter_operator"],
                            filter_value=row["filter_value"], target_value=row["target_value"])
                       for row in conditions],
        "nodes": [dict(node_id=row["node_id"], quest_id=row["quest_id"], node_type=row["node_type"],
                        reference=row["reference"], next_node=row["next_node"],
                        dialogue_phase=row["dialogue_phase"]) for row in nodes],
        "dialogue_slots": dialogue_slots, "dialogue": dialogue_lines}
    for name, payload in (("server_quests.json", server), ("client_quests.json", client)):
        (directory / name).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"server": str(directory / "server_quests.json"),
            "client": str(directory / "client_quests.json"), "accepted_quests": len(quests)}
